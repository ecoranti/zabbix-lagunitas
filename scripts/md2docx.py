#!/usr/bin/env python3
"""Convierte la documentación Markdown del proyecto a Word (.docx).

Uso:   .venv/bin/pip install python-docx
       .venv/bin/python scripts/md2docx.py [--portada docs/plantilla/portada-unrc.docx] \
           docs/guia-implementacion.md "docs/Guia de implementacion.docx"

Soporta lo que usan las guías: títulos, párrafos con **negrita**, `código` y
[enlaces](url), listas (con viñetas, numeradas y de verificación, con líneas de
continuación y bloques de código indentados), tablas, imágenes, bloques de
código, citas y separadores. Con --portada, el documento parte de esa plantilla
(portada institucional) y el contenido empieza en una página nueva.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
GRIS = RGBColor(0x55, 0x5F, 0x66)
TAM_TITULO = {1: 20, 2: 15, 3: 12.5}
MARCA_LISTA = re.compile(r"^(\s*)(- \[( |x)\]\s+|[-*]\s+|(\d+)\.\s+)(.*)")
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|<[^>]+>|\*[^*]+\*)")


def _sombrear(elemento, color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    elemento.append(shd)


def _inline(parrafo, texto: str) -> None:
    for parte in INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            parrafo.add_run(parte[2:-2]).bold = True
        elif parte.startswith("`") and parte.endswith("`"):
            r = parrafo.add_run(parte[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x8A, 0x1C, 0x3A)
        elif parte.startswith("[") and "](" in parte:
            txt, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", parte).groups()
            r = parrafo.add_run(txt)
            r.font.color.rgb = RGBColor(0x0A, 0x72, 0xB8)
            r.underline = True
            if url.startswith("http"):
                parrafo.add_run(f" ({url})").font.color.rgb = GRIS
        elif parte.startswith("<") and parte.endswith(">") and parte[1:5] == "http":
            parrafo.add_run(parte[1:-1]).font.color.rgb = RGBColor(0x0A, 0x72, 0xB8)
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            parrafo.add_run(parte[1:-1]).italic = True
        else:
            parrafo.add_run(parte)


def _tabla(doc, filas: list[list[str]]) -> None:
    encabezado, datos = filas[0], [f for f in filas[1:] if not all(re.fullmatch(r":?-{3,}:?", c) for c in f)]
    t = doc.add_table(rows=1, cols=len(encabezado))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(encabezado):
        celda = t.rows[0].cells[i]
        celda.text = ""
        _inline(celda.paragraphs[0], c)
        for r in celda.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _sombrear(celda._tc.get_or_add_tcPr(), "1F3A5F")
    for fila in datos:
        celdas = t.add_row().cells
        for i, c in enumerate(fila[:len(encabezado)]):
            celdas[i].text = ""
            _inline(celdas[i].paragraphs[0], c)
    for fila in t.rows:
        for celda in fila.cells:
            for p in celda.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()


def _codigo(doc, lineas: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_after = Pt(8)
    _sombrear(p._p.get_or_add_pPr(), "F2F4F6")
    r = p.add_run("\n".join(lineas))
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(8.5)


def _imagen(doc, alt: str, ruta: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ruta), width=Cm(16))
    if alt:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(alt)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GRIS


def _item(doc, sangria: str, marca: str, check: str | None, numero: str | None, texto: str) -> None:
    nivel = 1 if len(sangria) >= 2 else 0
    if check is not None or numero is not None:
        # Numeración literal: "List Number" de Word no reinicia entre listas.
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9 + 0.6 * nivel)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(("☐" if check == " " else "☑") if check is not None else f"{numero}.")
        p.add_run("\t")
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9 + 0.6 * nivel))
    else:
        p = doc.add_paragraph(style="List Bullet 2" if nivel else "List Bullet")
    _inline(p, texto)


def _titulo(doc, nivel: int, texto: str) -> None:
    h = doc.add_heading(level=min(nivel, 3))
    _inline(h, texto)
    # Formato directo (no de estilo) para no alterar los títulos de la portada.
    for r in h.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(TAM_TITULO[min(nivel, 3)])
        r.font.color.rgb = AZUL
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT


def convertir(origen: Path, destino: Path, portada: Path | None = None) -> None:
    doc = Document(str(portada)) if portada else Document()
    estilo = doc.styles["Normal"]
    if portada:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        # Word recalcula el índice (campo TOC de la plantilla) al abrir el documento.
        actualizar = OxmlElement("w:updateFields")
        actualizar.set(qn("w:val"), "true")
        doc.settings.element.append(actualizar)
    else:
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(10.5)
        for s in doc.sections:
            s.left_margin = s.right_margin = Cm(2.2)
            s.top_margin = s.bottom_margin = Cm(2)

    lineas = origen.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i]
        if linea.lstrip().startswith("```"):
            sangria = len(linea) - len(linea.lstrip())
            bloque = []
            i += 1
            while i < len(lineas) and not lineas[i].lstrip().startswith("```"):
                bloque.append(lineas[i][sangria:] if lineas[i][:sangria].isspace() else lineas[i].lstrip())
                i += 1
            _codigo(doc, bloque)
        elif m := re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", linea):
            _imagen(doc, m.group(1), origen.parent / m.group(2))
        elif linea.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                filas.append([c.strip() for c in lineas[i].strip().strip("|").split("|")])
                i += 1
            _tabla(doc, filas)
            continue
        elif m := re.match(r"^(#{1,4})\s+(.*)", linea):
            _titulo(doc, len(m.group(1)), m.group(2))
        elif re.match(r"^\s*-{3,}\s*$", linea):
            p = doc.add_paragraph()
            borde = OxmlElement("w:pBdr")
            abajo = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "C8CFD5")):
                abajo.set(qn(k), v)
            borde.append(abajo)
            p._p.get_or_add_pPr().append(borde)
        elif m := MARCA_LISTA.match(linea):
            texto = [m.group(5).strip()]
            # Líneas de continuación: indentadas, sin marca de lista ni bloque de código.
            while (i + 1 < len(lineas) and lineas[i + 1][:1].isspace() and lineas[i + 1].strip()
                   and not MARCA_LISTA.match(lineas[i + 1])
                   and not lineas[i + 1].lstrip().startswith(("```", "|", ">"))):
                i += 1
                texto.append(lineas[i].strip())
            _item(doc, m.group(1), m.group(2), m.group(3), m.group(4), " ".join(texto))
        elif linea.startswith(">"):
            texto = []
            while i < len(lineas) and lineas[i].startswith(">"):
                texto.append(lineas[i].lstrip("> ").rstrip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            _sombrear(p._p.get_or_add_pPr(), "FFF6E0")
            _inline(p, " ".join(texto))
            continue
        elif linea.strip():
            texto = [linea.strip()]
            while (i + 1 < len(lineas) and lineas[i + 1].strip()
                   and not re.match(r"^(#|\s*```|\||!\[|\s*[-*]\s|\s*\d+\.\s|>|\s*-{3,}\s*$)", lineas[i + 1])):
                i += 1
                texto.append(lineas[i].strip())
            _inline(doc.add_paragraph(), " ".join(texto))
        i += 1

    doc.core_properties.title = origen.stem
    doc.core_properties.author = "Elías Coranti — PPS UNRC"
    doc.save(destino)
    print(f"  {destino}")


if __name__ == "__main__":
    args = sys.argv[1:]
    portada = None
    if args[:1] == ["--portada"] and len(args) > 1:
        portada, args = Path(args[1]), args[2:]
    if len(args) != 2:
        sys.exit(__doc__)
    convertir(Path(args[0]), Path(args[1]), portada)
